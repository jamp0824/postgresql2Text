"""
문서 생성 모듈 — Markdown / Word(.docx) 포맷 문서 출력
"""

from __future__ import annotations

import re
import logging
from datetime import datetime
from pathlib import Path

from pg2text.config import settings
from pg2text.nl_processor import GeneratedDocument

logger = logging.getLogger(__name__)


# ─── Markdown 빌더 ────────────────────────────────────────────────────────────


class MarkdownBuilder:
    """Markdown 파일 저장"""

    def save(
        self,
        document: GeneratedDocument,
        output_path: Path | None = None,
    ) -> Path:
        """
        문서를 .md 파일로 저장.

        Args:
            document: 저장할 문서
            output_path: 저장 경로 (None이면 자동 생성)
        Returns:
            저장된 파일 경로
        """
        if output_path is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_title = re.sub(r"[^\w가-힣\-]", "_", document.title)[:40]
            output_path = settings.output_dir / f"{safe_title}_{timestamp}.md"

        output_path.parent.mkdir(parents=True, exist_ok=True)

        # 메타데이터 헤더 추가
        meta_header = f"""---
title: "{document.title}"
generated_at: "{datetime.now().isoformat()}"
---

"""
        full_content = meta_header + document.markdown_content
        output_path.write_text(full_content, encoding="utf-8")
        logger.info("Markdown 저장: %s", output_path)
        return output_path


# ─── Word(.docx) 빌더 ────────────────────────────────────────────────────────


class DocxBuilder:
    """Word 문서(.docx) 생성"""

    def save(
        self,
        document: GeneratedDocument,
        output_path: Path | None = None,
    ) -> Path:
        """
        문서를 .docx 파일로 저장.

        Args:
            document: 저장할 문서
            output_path: 저장 경로 (None이면 자동 생성)
        Returns:
            저장된 파일 경로
        """
        try:
            from docx import Document as DocxDocument
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml.ns import qn
        except ImportError:
            raise ImportError(
                "python-docx가 설치되지 않았습니다. "
                "`pip install python-docx` 로 설치하세요."
            )

        if output_path is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_title = re.sub(r"[^\w가-힣\-]", "_", document.title)[:40]
            output_path = settings.output_dir / f"{safe_title}_{timestamp}.docx"

        output_path.parent.mkdir(parents=True, exist_ok=True)

        doc = DocxDocument()
        self._setup_styles(doc)
        self._parse_markdown_to_docx(doc, document.markdown_content)

        # 푸터: 생성 일시
        section = doc.sections[0]
        footer = section.footer
        footer_para = footer.paragraphs[0]
        footer_para.text = f"생성일시: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        doc.save(str(output_path))
        logger.info("Word 문서 저장: %s", output_path)
        return output_path

    def _setup_styles(self, doc) -> None:
        """문서 기본 스타일 설정"""
        from docx.shared import Pt
        from docx.oxml.ns import qn
        import lxml.etree as etree

        # 기본 폰트 설정 (한글 지원)
        style = doc.styles["Normal"]
        font = style.font
        font.name = "맑은 고딕"
        font.size = Pt(10.5)

    def _parse_markdown_to_docx(self, doc, markdown: str) -> None:
        """Markdown을 Word 문서 요소로 변환"""
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        lines = markdown.split("\n")
        in_table = False
        table_rows: list[list[str]] = []
        in_code_block = False
        code_lines: list[str] = []

        for line in lines:
            # 코드 블록 처리
            if line.strip().startswith("```"):
                if in_code_block:
                    # 코드 블록 종료 → Word에 삽입
                    if code_lines:
                        para = doc.add_paragraph("\n".join(code_lines))
                        para.style = "No Spacing"
                        for run in para.runs:
                            run.font.name = "Courier New"
                            run.font.size = Pt(9)
                    code_lines = []
                    in_code_block = False
                else:
                    in_code_block = True
                continue

            if in_code_block:
                code_lines.append(line)
                continue

            # 테이블 처리
            if line.strip().startswith("|"):
                # 구분선 행 스킵
                if re.match(r"^\|\s*[-:]+\s*(\|\s*[-:]+\s*)*\|?\s*$", line):
                    continue
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                table_rows.append(cells)
                in_table = True
                continue
            elif in_table and table_rows:
                # 테이블 끝 → Word 테이블 생성
                self._add_table(doc, table_rows)
                table_rows = []
                in_table = False

            # 헤딩
            if line.startswith("# "):
                doc.add_heading(line[2:].strip(), level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=3)
            elif line.startswith("#### "):
                doc.add_heading(line[5:].strip(), level=4)
            # 수평선
            elif line.strip() in ("---", "***", "___"):
                doc.add_paragraph("─" * 40)
            # 불릿 리스트
            elif line.startswith("- ") or line.startswith("* "):
                para = doc.add_paragraph(
                    self._strip_inline_md(line[2:].strip()),
                    style="List Bullet",
                )
            # 번호 리스트
            elif re.match(r"^\d+\.\s", line):
                content = re.sub(r"^\d+\.\s", "", line)
                doc.add_paragraph(
                    self._strip_inline_md(content.strip()),
                    style="List Number",
                )
            # YAML 메타데이터 헤더 스킵
            elif line.strip() in ("---",) or line.strip().startswith("title:") \
                    or line.strip().startswith("generated_at:"):
                continue
            # 빈 줄
            elif not line.strip():
                doc.add_paragraph("")
            # 일반 텍스트
            else:
                para = doc.add_paragraph()
                self._add_inline_runs(para, line)

        # 남은 테이블 처리
        if in_table and table_rows:
            self._add_table(doc, table_rows)

    def _add_table(self, doc, table_rows: list[list[str]]) -> None:
        """Word 테이블 추가"""
        from docx.shared import RGBColor, Pt
        from docx.oxml.ns import qn

        if not table_rows:
            return

        num_cols = max(len(row) for row in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=num_cols)
        table.style = "Table Grid"

        for r_idx, row in enumerate(table_rows):
            for c_idx, cell_text in enumerate(row):
                if c_idx >= num_cols:
                    break
                cell = table.rows[r_idx].cells[c_idx]
                cell.text = cell_text
                # 헤더 행 굵게
                if r_idx == 0:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True

        doc.add_paragraph("")  # 테이블 후 여백

    def _strip_inline_md(self, text: str) -> str:
        """인라인 Markdown 마크업 제거 (굵게, 기울임 등)"""
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        text = re.sub(r"\*(.+?)\*", r"\1", text)
        text = re.sub(r"`(.+?)`", r"\1", text)
        text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)
        return text

    def _add_inline_runs(self, para, text: str) -> None:
        """인라인 Markdown(굵게/이탤릭/코드)을 Word run으로 변환"""
        from docx.shared import Pt

        pattern = re.compile(r"(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`|(.+?)(?=\*\*|\*|`|$))")
        pos = 0
        while pos < len(text):
            bold_match = re.match(r"\*\*(.+?)\*\*", text[pos:])
            italic_match = re.match(r"\*(.+?)\*", text[pos:])
            code_match = re.match(r"`(.+?)`", text[pos:])

            if bold_match:
                run = para.add_run(bold_match.group(1))
                run.bold = True
                pos += len(bold_match.group(0))
            elif italic_match:
                run = para.add_run(italic_match.group(1))
                run.italic = True
                pos += len(italic_match.group(0))
            elif code_match:
                run = para.add_run(code_match.group(1))
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                pos += len(code_match.group(0))
            else:
                # 다음 마크업까지 일반 텍스트
                next_marker = len(text)
                for marker in ["**", "*", "`"]:
                    idx = text.find(marker, pos)
                    if idx != -1 and idx < next_marker:
                        next_marker = idx
                para.add_run(text[pos:next_marker])
                pos = next_marker


# ─── 문서 빌더 팩토리 ─────────────────────────────────────────────────────────


class DocumentBuilder:
    """
    통합 문서 빌더 — 형식에 따라 적절한 빌더 선택
    """

    def __init__(self):
        self._md_builder = MarkdownBuilder()
        self._docx_builder = DocxBuilder()

    def save(
        self,
        document: GeneratedDocument,
        format: str = "md",
        output_path: Path | None = None,
    ) -> Path:
        """
        문서를 지정된 형식으로 저장.

        Args:
            document: 저장할 문서
            format: "md" 또는 "docx"
            output_path: 저장 경로 (None이면 자동 생성)
        Returns:
            저장된 파일 경로
        """
        if format.lower() in ("md", "markdown"):
            return self._md_builder.save(document, output_path)
        elif format.lower() in ("docx", "word"):
            return self._docx_builder.save(document, output_path)
        else:
            raise ValueError(f"지원하지 않는 형식: {format}. 'md' 또는 'docx'를 사용하세요.")

    def save_all(
        self, document: GeneratedDocument, output_dir: Path | None = None
    ) -> dict[str, Path]:
        """모든 형식으로 저장"""
        base_dir = output_dir or settings.output_dir
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_title = re.sub(r"[^\w가-힣\-]", "_", document.title)[:40]
        base_name = f"{safe_title}_{timestamp}"

        results = {}
        results["md"] = self._md_builder.save(
            document, base_dir / f"{base_name}.md"
        )
        results["docx"] = self._docx_builder.save(
            document, base_dir / f"{base_name}.docx"
        )
        return results
